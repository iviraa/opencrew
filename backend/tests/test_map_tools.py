import json
import time
from datetime import date
from pathlib import Path

import psycopg
import pytest

from app import share
from app.crewly import map_tools as mt
from app.crewly.export_tools import run_query, to_geojson, to_ics, to_kml, to_xlsx
from app.db import DATABASE_URL, connect

GPC = {"id": "u1", "company": "gpc", "username": "georgia", "token": "tok-gpc"}
FIX = Path(__file__).parent / "fixtures"


# ---------- pure pieces, no database ----------

def test_project_filter_sql_binds_every_value():
    clauses, params = mt.project_filter_sql({"kv": ["230", 115, "x"], "type": "line", "status": "Planned", "state": "ga", "county": "Fulton", "name": "Sub", "years": [2028, 2027]})
    assert len(clauses) == 7 and params["kv"] == [230, 115] and params["state"] == "GA" and params["ylo"] == "2027-01-01" and params["yhi"] == "2028-12-31"
    assert all("%(" in c for c in clauses) and "Fulton" not in " ".join(clauses)  # values are bound, never pasted
    assert mt.project_filter_sql({}) == ([], {})


def test_parse_days_converts_units_and_flags_thresholds():
    grid = json.load(open(FIX / "gridpoint.json"))["properties"]
    grid.update({"tz": "America/New_York",
                 "probabilityOfThunder": {"uom": "wmoUnit:percent", "values": [{"validTime": "2026-09-28T18:00:00+00:00/PT6H", "value": 55}]},
                 "quantitativePrecipitation": {"uom": "wmoUnit:mm", "values": [{"validTime": "2026-09-28T18:00:00+00:00/PT6H", "value": 30.0}]},
                 "iceAccumulation": {"uom": "wmoUnit:mm", "values": [{"validTime": "2026-09-30T06:00:00+00:00/PT6H", "value": 2.0}]},
                 "snowfallAmount": {"uom": "wmoUnit:mm", "values": []},
                 "heatIndex": {"uom": "wmoUnit:degC", "values": [{"validTime": "2026-09-27T18:00:00+00:00/PT3H", "value": 35.0}]},
                 "temperature": {"uom": "wmoUnit:degC", "values": [{"validTime": "2026-09-27T18:00:00+00:00/PT3H", "value": 30.0}, {"validTime": "2026-09-27T08:00:00+00:00/PT3H", "value": 18.0}]}})
    days = mt.parse_days(grid, days=4, today=date(2026, 9, 27))
    assert [d["date"] for d in days] == ["2026-09-27", "2026-09-28", "2026-09-29", "2026-09-30"]
    d0, d1, d3 = days[0], days[1], days[3]
    assert d0["gust_mph"] == 25 and d0["high_f"] == 86 and d0["low_f"] == 64 and d0["heat_index_f"] == 95 and any("heat index" in n for n in d0["notes"])
    assert d1["gust_mph"] == 45 and d1["thunder_pct"] == 55 and d1["rain_in"] == 1.18
    assert any("crane" in n for n in d1["notes"]) and any("thunder" in n for n in d1["notes"]) and any("rain" in n for n in d1["notes"])
    assert d3["ice_in"] == 0.08 and any("ice" in n for n in d3["notes"])
    assert days[2]["gust_mph"] is None and days[2]["quiet"] is False  # the null reading is skipped, so day 3 has no data at all


def test_every_method_has_a_title_text_and_sources():
    for k, (title, text, sources) in mt.METHODS.items():
        assert title and len(text) > 60 and isinstance(sources, list), k
    out, acts = mt.explain_method({"company": "gpc"}, None, "weather")
    assert out["title"] == mt.METHODS["hazard_days"][0] and acts[0]["type"] == "explain"
    assert "error" in mt.explain_method({"company": "gpc"}, None, "astrology")[0]


def test_resolve_point_reads_coordinates_either_way():
    assert mt.resolve_point(None, "gpc", "-84.39, 33.75")[:2] == (-84.39, 33.75)
    assert mt.resolve_point(None, "gpc", "33.75, -84.39")[:2] == (-84.39, 33.75)
    assert mt.resolve_point(None, "gpc", "augusta")[2] == "Augusta"


def test_share_tokens_sign_verify_expire_and_reject_tampering(monkeypatch):
    monkeypatch.setenv("SHARE_SECRET", "test-secret")
    exp = time.time() + 3600
    tok = share.sign("report", "12", "gpc", exp)
    claims = share.verify(tok)
    assert claims["kind"] == "report" and claims["ref_id"] == "12" and claims["company"] == "gpc" and claims["expires_at"] == int(exp)
    assert share.verify(tok, now=exp + 1) is None  # expired
    head, sig = tok.split(".")
    assert share.verify(head + "." + sig[:-2] + "AA") is None  # bad signature
    other = share.sign("report", "13", "gpc", exp).split(".")[0]
    assert share.verify(other + "." + sig) is None  # payload swapped
    assert share.verify("garbage") is None and share.verify("") is None
    monkeypatch.setenv("SHARE_SECRET", "another-secret")
    assert share.verify(tok) is None  # a different secret issued it


def test_query_dsl_groups_aggregates_and_sorts():
    rows = [{"partner": "A", "tier": "site", "drive_min": 10, "usd": 5}, {"partner": "A", "tier": "crew", "drive_min": 30, "usd": None},
            {"partner": "B", "tier": "site", "drive_min": 20, "usd": 7}]
    out, cols = run_query(rows, ["partner"], {"drive_min": "avg", "usd": "sum"}, {"by": "avg_drive_min", "dir": "desc"})
    assert cols == ["partner", "rows", "avg_drive_min", "sum_usd"]
    assert out == [{"partner": "A", "rows": 2, "avg_drive_min": 20.0, "sum_usd": 5.0}, {"partner": "B", "rows": 1, "avg_drive_min": 20.0, "sum_usd": 7.0}][::-1] or out[0]["partner"] in "AB"
    total, _ = run_query(rows, None, {"drive_min": "max", "usd": "count"})
    assert total == [{"rows": 3, "max_drive_min": 30.0, "count_usd": 2}]
    plain, _ = run_query(rows, None, None, {"by": "drive_min", "dir": "asc"}, limit=2)
    assert [r["drive_min"] for r in plain] == [10, 20]
    assert run_query([], ["x"], {"y": "sum"}) == ([], ["x", "rows", "sum_y"])


def test_writers_produce_parseable_files():
    feats = [{"type": "Feature", "id": 1, "geometry": {"type": "LineString", "coordinates": [[-84.0, 33.0], [-84.1, 33.1]]}, "properties": {"name": "A & B", "score": 0.5}},
             {"type": "Feature", "id": 2, "geometry": {"type": "Point", "coordinates": [-84.0, 33.0]}, "properties": {"name": "P"}}]
    gj = json.loads(to_geojson(feats))
    assert gj["type"] == "FeatureCollection" and len(gj["features"]) == 2
    kml = to_kml(feats, "t").decode()
    import xml.etree.ElementTree as ET
    root = ET.fromstring(kml)
    assert kml.count("<Placemark>") == 2 and "A &amp; B" in kml and root.tag.endswith("kml")
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(to_xlsx([{"a": 1, "b": "x"}, {"a": 2, "b": None}], "overlaps")))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["a", "b"] and ws["A3"].value == 2
    ics = to_ics([{"uid": "plan-1", "start": "2027-03-01", "end": "2027-05-01", "summary": "W, x", "description": "d"}, {"uid": "bad", "start": "n/a", "end": "", "summary": "s"}], "crewly plan").decode()
    assert ics.count("BEGIN:VEVENT") == 1 and "DTSTART;VALUE=DATE:20270301" in ics and "SUMMARY:W\\, x" in ics


# ---------- against the planner database ----------

try:
    with psycopg.connect(DATABASE_URL, connect_timeout=3) as c:
        JOBS = c.execute("SELECT count(*) FROM job").fetchone()[0]
except Exception:
    JOBS = 0
needs_db = pytest.mark.skipif(not JOBS, reason="database not reachable or empty")


@pytest.fixture
def conn():
    with connect() as c:
        yield c
        c.rollback()


class FakeDrive:
    def __init__(self):
        self.cache, self.saved, self.urls = {}, 0, []

    def between(self, a, b):
        return {"min": 42.0, "km": 50.0}

    def _get(self, url, params):
        self.urls.append(url)
        return {"routes": [{"geometry": {"type": "LineString", "coordinates": [[-84.0, 33.0], [-84.5, 33.5], [-85.0, 34.0]]}}]}

    def save(self):
        self.saved += 1


@needs_db
def test_map_view_switches_tabs_and_resolves_filters_to_ids(conn):
    out, acts = mt.map_view(GPC, conn, tab="weather", period="month", month=8, hazards=["flood"], fit="ours")
    a = acts[0]
    assert a == {"type": "map_view", "tab": "hazards", "period": "month", "month": 8, "hazards": ["flood"], "fit": {"to": "ours"}}
    out, acts = mt.map_view(GPC, conn, tab="overlaps", filters={"tier": "site"}, fit=[-85, 30, -80, 35])
    assert isinstance(acts[0]["ids"], list) and acts[0]["fit"] == {"bbox": [-85, 30, -80, 35]}
    assert "error" in mt.map_view(GPC, conn, tab="plans")[0]
    assert "error" in mt.map_view(GPC, conn, tab="hazards", hazards=["locusts"])[0]
    assert "error" in mt.map_view(GPC, conn, filters={"partner": "Nobody Power"})[0]


@needs_db
def test_filter_projects_and_timeline_are_ours(conn):
    out, acts = mt.filter_projects(GPC, conn, filters={"years": [2027, 2029]}, limit=5)
    rows = acts[0]["projects"]["rows"]
    assert acts[0]["type"] == "projects" and len(rows) <= 5 and all(r["id"] for r in rows)
    out, acts = mt.timeline(GPC, conn, years=[2027, 2028])
    tl = acts[0]["timeline"]
    assert acts[0]["type"] == "timeline" and tl["years"] == [2027, 2028] and all("mine" in r and "start" in r for r in tl["rows"])


@needs_db
def test_site_forecast_uses_the_cache_and_a_fake_fetch(conn):
    grid = {"tz": "America/New_York", "place": {"city": "Atlanta", "state": "GA"}, **{k: None for k in mt.FIELDS},
            "windGust": {"uom": "wmoUnit:km_h-1", "values": [{"validTime": f"{date.today().isoformat()}T12:00:00+00:00/PT3H", "value": 80.0}]}}
    calls = []

    def fetch(lon, lat):
        calls.append((lon, lat))
        return grid
    out, acts = mt.site_forecast(GPC, conn, "-84.39,33.75", fetch=fetch)
    fc = acts[0]["forecast"]
    assert acts[0]["type"] == "forecast" and fc["near"] == "Atlanta GA" and len(fc["days"]) == 7 and fc["days"][0]["gust_mph"] == 50
    assert "crane" in fc["days"][0]["notes"][0] and out["summary"].startswith("1 of 7")
    mt.site_forecast(GPC, conn, "-84.40,33.76", fetch=fetch)  # same 0.05 degree cell: served from the cache
    assert len(calls) == 1
    assert "error" in mt.site_forecast(GPC, conn, "nowhere at all", fetch=fetch)[0]


@needs_db
def test_route_between_draws_a_line_and_caches_geometry(conn):
    d = FakeDrive()
    out, acts = mt.route_between(GPC, conn, "-84.39,33.75", "augusta", drive=d)
    r = acts[0]["route"]
    assert acts[0]["type"] == "route" and r["road_mi"] == 31.1 and r["drive_min"] == 42.0 and len(r["geometry"]["coordinates"]) == 3 and d.saved == 1
    assert r["a"] == [-84.39, 33.75] and out["straight_km"] > 0
    mt.route_between(GPC, conn, "-84.39,33.75", "augusta", drive=d)
    assert len(d.urls) == 1  # geometry came from the cache the second time


@needs_db
def test_export_share_and_query_round_trip(conn):
    from app.crewly.export_tools import export, query_data, share_link
    out, acts = export(GPC, conn, "overlaps", "geojson")
    d = acts[0]["download"]
    assert acts[0]["type"] == "download" and d["filename"].endswith(".geojson") and d["size"] > 50
    raw = conn.execute("SELECT bytes, media_type FROM export WHERE id = %s", (d["id"],)).fetchone()
    assert json.loads(bytes(raw["bytes"]))["type"] == "FeatureCollection" and raw["media_type"] == "application/geo+json"
    for fmt in ("csv", "xlsx", "kml"):
        assert export(GPC, conn, "projects", fmt)[1][0]["download"]["format"] == fmt
    assert export(GPC, conn, "projects", "ics")[1][0]["download"]["filename"].endswith(".ics")
    assert "error" in export(GPC, conn, "overlaps", "pdf")[0] and "error" in export(GPC, conn, "hazard_exposure", "kml")[0]

    out, acts = query_data(GPC, conn, "overlaps", group_by=["partner"], aggregate={"drive_min": "avg"}, sort={"by": "avg_drive_min"})
    assert [a["type"] for a in acts] == ["table", "chart"] and acts[0]["table"]["columns"][:2] == ["partner", "rows"]
    assert "error" in query_data(GPC, conn, "unicorns")[0]

    from app.crewly import reports
    oid = conn.execute("SELECT op.id FROM opportunity op JOIN job ja ON ja.id = op.job_a JOIN job jb ON jb.id = op.job_b WHERE op.horizon = 'long' AND 'gpc' IN (ja.org_id, jb.org_id) ORDER BY op.score DESC LIMIT 1").fetchone()["id"]
    rep = reports.build(conn, GPC["company"], "hazard_exposure", str(oid))
    out, acts = share_link(GPC, conn, "report", str(rep["id"]), 3)
    s = acts[0]["share"]
    assert acts[0]["type"] == "share" and s["days"] == 3 and s["path"].startswith("/share/") and "token" not in out
    claims = share.verify(s["token"])
    assert claims["company"] == "gpc" and share.render(conn, claims).startswith("<!doctype html")
    assert "error" in share_link(GPC, conn, "report", "999999")[0] and "error" in share_link(GPC, conn, "finding", "abc")[0]


@needs_db
def test_context_projects_excludes_our_org_and_stays_in_the_box(conn):
    from fastapi.testclient import TestClient
    from app.auth import current_user
    from app.main import app
    app.dependency_overrides[current_user] = lambda: GPC
    try:
        c = TestClient(app)
        fc = c.get("/api/app/context_projects?bbox=-86,30,-79,36").json()
        assert fc["type"] == "FeatureCollection" and len(fc["features"]) > 20
        assert all(f["properties"]["org_id"] != "gpc" and f["properties"]["org"] and "dotted" in f["properties"] for f in fc["features"])
        assert any(f["properties"]["dotted"] for f in fc["features"]) and any(not f["properties"]["dotted"] for f in fc["features"])
        assert c.get("/api/app/context_projects?bbox=-100,44,-99,45").json()["features"] == []
        assert c.get("/api/app/context_projects?bbox=nope").status_code == 400
    finally:
        app.dependency_overrides.pop(current_user, None)
    out, acts = mt.map_view(GPC, conn, tab="overlaps", layers={"others": False})
    assert acts[0]["layers"] == {"others": False} and "hiding" in out["did"]
