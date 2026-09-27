"""Crewly abilities that hand data on: files in the formats planners use, expiring share links, and grouped queries over the catalogue."""
import io
import json
import re
from datetime import date, datetime, timezone
from xml.sax.saxutils import escape

from app import share
from app.companies import name
from app.crewly import charts

EXPORT_KINDS = ("overlaps", "projects", "plan", "findings", "hazard_exposure")
FORMATS = {"geojson": ("application/geo+json", "geojson"), "kml": ("application/vnd.google-earth.kml+xml", "kml"), "csv": ("text/csv", "csv"),
           "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"), "ics": ("text/calendar", "ics")}
AGGS = ("sum", "avg", "min", "max", "count")


def _bind(ctx, fn):
    return lambda conn, **kw: fn(ctx, conn, **kw)


def _int(v, default=None):
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


# ---------- rows and geometries per kind ----------

def _rows(conn, ctx, kind, filters):
    company = ctx["company"]
    if kind == "findings":
        from app.scenario import experiments
        return [{"id": f["id"], "title": f["title"], "kind": f["kind"], "starred": f.get("starred"), "created_at": str(f.get("created_at") or "")[:19],
                 **{f"delta_{d['metric']}": d.get("delta") for d in (f.get("deltas") or [])[:8]}} for f in experiments.listing(conn, company, bool(filters.get("starred")), 200)]
    if kind == "plan":
        _, rows = charts.plan_totals(conn, company, filters)
        return rows
    return charts.table(conn, company, {"hazard_exposure": "hazard_exposure"}.get(kind, kind), filters, ctx)["rows"]


def _features(conn, ctx, kind, filters, rows):
    """GeoJSON features for kinds that have places: overlaps (link lines), projects (works), plan (its overlaps' links)."""
    company = ctx["company"]
    if kind == "projects":
        ids = [r["project"] for r in rows]
        got = conn.execute("SELECT id, ST_AsGeoJSON(geom::geometry, 5)::json AS g FROM job WHERE id = ANY(%s)", (ids,)).fetchall()
        geoms = {g["id"]: g["g"] for g in got}
        return [{"type": "Feature", "id": r["project"], "geometry": geoms.get(r["project"]), "properties": {k: v for k, v in r.items()}} for r in rows if geoms.get(r["project"])]
    key = "overlap"
    ids = [r[key] for r in rows if r.get(key) is not None]
    if not ids:
        return []
    got = conn.execute("SELECT op.id, ST_AsGeoJSON(op.link::geometry, 5)::json AS g FROM opportunity op JOIN job ja ON ja.id = op.job_a JOIN job jb ON jb.id = op.job_b "
                       "WHERE op.id = ANY(%s) AND %s IN (ja.org_id, jb.org_id)", (ids, company)).fetchall()
    geoms = {g["id"]: g["g"] for g in got}
    return [{"type": "Feature", "id": r[key], "geometry": geoms.get(r[key]), "properties": {k: v for k, v in r.items()}} for r in rows if geoms.get(r[key])]


# ---------- writers (pure: rows/features in, bytes out) ----------

def to_geojson(features):
    return json.dumps({"type": "FeatureCollection", "features": features}, default=str).encode()


def _kml_geom(g):
    if not g:
        return ""
    t, c = g["type"], g["coordinates"]
    if t == "Point":
        return f"<Point><coordinates>{c[0]},{c[1]}</coordinates></Point>"
    if t == "LineString":
        return "<LineString><coordinates>" + " ".join(f"{x},{y}" for x, y in c) + "</coordinates></LineString>"
    if t == "MultiLineString":
        return "<MultiGeometry>" + "".join(_kml_geom({"type": "LineString", "coordinates": ln}) for ln in c) + "</MultiGeometry>"
    if t == "Polygon":
        return "<Polygon><outerBoundaryIs><LinearRing><coordinates>" + " ".join(f"{x},{y}" for x, y in c[0]) + "</coordinates></LinearRing></outerBoundaryIs></Polygon>"
    return ""


def to_kml(features, title):
    marks = []
    for f in features:
        p = f.get("properties") or {}
        label = str(p.get("name") or p.get("ours") or f.get("id") or "")
        desc = "<br/>".join(f"{escape(str(k))}: {escape(str(v))}" for k, v in p.items() if v not in (None, "") and k not in ("name",))
        marks.append(f"<Placemark><name>{escape(label)}</name><description><![CDATA[{desc}]]></description>{_kml_geom(f.get('geometry'))}</Placemark>")
    return (f'<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>{escape(title)}</name>'
            + "".join(marks) + "</Document></kml>").encode()


def to_xlsx(rows, title):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[^A-Za-z0-9 _-]", "", title)[:28] or "data"
    cols = list(rows[0].keys()) if rows else []
    ws.append(cols)
    for r in rows:
        ws.append([_cell(r.get(c)) for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, len(c) + 2, *(len(str(r.get(c, ""))) for r in rows[:200])))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v, default=str)
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return v


def _ics_date(s):
    d = str(s)[:10]
    return d.replace("-", "") if re.fullmatch(r"\d{4}-\d{2}-\d{2}", d) else None


def to_ics(events, calname):
    """All-day events; each is {uid, start, end, summary, description}."""
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//crewly//coordination//EN", f"X-WR-CALNAME:{_ics_text(calname)}"]
    for e in events:
        s, t = _ics_date(e["start"]), _ics_date(e["end"])
        if not s or not t:
            continue
        lines += ["BEGIN:VEVENT", f"UID:{e['uid']}@crewly", f"DTSTAMP:{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}", f"DTSTART;VALUE=DATE:{s}", f"DTEND;VALUE=DATE:{t}",
                  f"SUMMARY:{_ics_text(e['summary'])}", f"DESCRIPTION:{_ics_text(e.get('description') or '')}", "END:VEVENT"]
    lines.append("END:VCALENDAR")
    return ("\r\n".join(lines) + "\r\n").encode()


def _ics_text(s):
    return str(s).replace("\\", "\\\\").replace("\n", "\\n").replace(",", "\\,").replace(";", "\\;")[:400]


def _events(kind, rows, company):
    if kind == "plan":
        return [{"uid": f"plan-{r['overlap']}", "start": r["months"].split(" to ")[0] + "-01", "end": _month_end(r["months"].split(" to ")[1]),
                 "summary": f"Target window: {r['ours']} with {r['partner']}", "description": f"overlap #{r['overlap']}, {r['verdict']}, savings {r['savings_low']} to {r['savings_high']} USD"} for r in rows]
    if kind == "projects":
        return [{"uid": f"project-{r['project']}", "start": r["start"], "end": r["end"], "summary": r["name"], "description": f"{r.get('kv') or ''} kV {r.get('type') or ''}, in service {r.get('in_service')}"} for r in rows]
    return []


def _month_end(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    return f"{y + (m == 12)}-{m % 12 + 1:02d}-01"


def build_file(conn, ctx, kind, fmt, filters):
    """(filename, media type, bytes) for one export, or raises ValueError with a plain reason."""
    if kind not in EXPORT_KINDS:
        raise ValueError(f"kind must be one of {', '.join(EXPORT_KINDS)}")
    if fmt not in FORMATS:
        raise ValueError(f"format must be one of {', '.join(FORMATS)}")
    rows = _rows(conn, ctx, kind, filters or {})
    stamp = date.today().isoformat()
    base = f"crewly-{kind}-{stamp}"
    media, ext = FORMATS[fmt]
    if fmt == "csv":
        return f"{base}.csv", media, charts.to_csv(rows).encode()
    if fmt == "xlsx":
        return f"{base}.xlsx", media, to_xlsx(rows, kind)
    if fmt == "ics":
        events = _events(kind, rows, ctx["company"])
        if not events:
            raise ValueError("calendar files need plan target windows or project build windows; export the plan or projects as ics")
        return f"{base}.ics", media, to_ics(events, f"crewly {kind}")
    feats = _features(conn, ctx, kind, filters or {}, rows)
    if not feats:
        raise ValueError(f"{kind} has no map shapes to export; try csv or xlsx")
    if fmt == "kml":
        return f"{base}.kml", media, to_kml(feats, f"crewly {kind} for {name(ctx['company'])}")
    return f"{base}.geojson", media, to_geojson(feats)


def export(ctx, conn, kind, format="geojson", filters=None):
    try:
        filename, media, data = build_file(conn, ctx, kind, format, filters)
    except ValueError as e:
        return {"error": str(e)}, []
    conn.execute(share.TABLE_SQL)
    row = conn.execute("INSERT INTO export (company_id, kind, format, filename, media_type, bytes) VALUES (%s, %s, %s, %s, %s, %s) RETURNING id",
                       (ctx["company"], kind, format, filename, media, data)).fetchone()
    card = {"id": row["id"], "kind": kind, "format": format, "filename": filename, "size": len(data), "filters": {k: v for k, v in (filters or {}).items() if v not in (None, "", [])}}
    return {**card, "next_step": "the download card in the chat saves the file"}, [{"type": "download", "download": card}]


# ---------- share links ----------

def share_link(ctx, conn, kind, id, expires_days=7):
    ref = str(id or "").lstrip("#")
    if kind == "plan":
        ref = ref or "quarter"
    elif not ref.isdigit():
        return {"error": f"a {kind} share needs its numeric id"}, []
    if kind == "report":
        from app.crewly import reports
        row = reports.get(conn, ref, ctx["company"])
        if not row:
            return {"error": f"report {ref} is not ours or does not exist"}, []
        title = row["title"]
    elif kind == "finding":
        from app.scenario import experiments
        f = experiments.get(conn, int(ref), ctx["company"])
        if not f:
            return {"error": f"finding {ref} is not ours or does not exist"}, []
        title = f["title"]
    elif kind == "plan":
        from app.planner import build, store
        row = store.latest(conn, ctx["company"], ref if ref in build.HORIZONS else "quarter")
        if not row:
            return {"error": "no plan yet; ask for one first"}, []
        title = f"Coordination plan, {row['horizon']} v{row['version']}"
    else:
        return {"error": f"kind must be one of {', '.join(share.KINDS)}"}, []
    try:
        s = share.create(conn, ctx["company"], kind, ref, title, expires_days)
    except ValueError as e:
        return {"error": str(e)}, []
    return {**{k: v for k, v in s.items() if k != "token"}, "note": "anyone with the link can read it until it expires or you revoke it"}, [{"type": "share", "share": s}]


# ---------- grouped queries over the catalogue ----------

def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def run_query(rows, group_by=None, aggregate=None, sort=None, limit=200):
    """Deterministic group / aggregate / sort over dict rows; returns (rows, columns)."""
    group_by = [g for g in (group_by or []) if g]
    aggregate = {k: v for k, v in (aggregate or {}).items() if v in AGGS}
    if group_by:
        buckets = {}
        for r in rows:
            key = tuple(str(r.get(g, "")) for g in group_by)
            buckets.setdefault(key, []).append(r)
        out = []
        for key, rs in buckets.items():
            row = dict(zip(group_by, key))
            row["rows"] = len(rs)
            for field, fn in aggregate.items():
                vals = [x for x in (_num(r.get(field)) for r in rs) if x is not None]
                row[f"{fn}_{field}"] = (len(vals) if fn == "count" else None if not vals else round(sum(vals) if fn == "sum" else sum(vals) / len(vals) if fn == "avg" else min(vals) if fn == "min" else max(vals), 2))
            out.append(row)
        rows = out
    elif aggregate:
        row = {"rows": len(rows)}
        for field, fn in aggregate.items():
            vals = [x for x in (_num(r.get(field)) for r in rows) if x is not None]
            row[f"{fn}_{field}"] = (len(vals) if fn == "count" else None if not vals else round(sum(vals) if fn == "sum" else sum(vals) / len(vals) if fn == "avg" else min(vals) if fn == "min" else max(vals), 2))
        rows = [row]
    if sort and sort.get("by"):
        by, desc = sort["by"], str(sort.get("dir", "desc")).lower().startswith("d")
        rows = sorted(rows, key=lambda r: (r.get(by) is None, r.get(by) if isinstance(r.get(by), (int, float)) else str(r.get(by) or "")), reverse=desc)
    rows = rows[: max(1, min(_int(limit, 200), 500))]
    cols = list(rows[0].keys()) if rows else (group_by + ["rows"] + [f"{fn}_{f}" for f, fn in aggregate.items()])
    return rows, cols


def query_data(ctx, conn, dataset, filters=None, group_by=None, aggregate=None, sort=None, limit=200):
    try:
        base = charts.table(conn, ctx["company"], dataset, filters or {}, ctx)
    except ValueError as e:
        return {"error": str(e)}, []
    rows, cols = run_query(base["rows"], group_by, aggregate, sort, limit)
    table = {"dataset": dataset, "filters": {**base["filters"], **({"group_by": group_by} if group_by else {}), **({"aggregate": ", ".join(f"{fn} {f}" for f, fn in aggregate.items())} if aggregate else {})},  # chips read as text
             "columns": cols, "rows": rows, "count": len(rows)}
    actions = [{"type": "table", "table": table}]
    if group_by and aggregate and rows:
        first = next(iter(aggregate.items()))
        col = f"{first[1]}_{first[0]}"
        chart = {"title": f"{col.replace('_', ' ')} by {', '.join(group_by)}", "kind": "bar", "dataset": dataset, "options": {"group_by": group_by},
                 "x": [" / ".join(str(r.get(g)) for g in group_by) for r in rows[:24]], "series": [{"name": col.replace("_", " "), "values": [r.get(col) or 0 for r in rows[:24]], "color": "#5b2bb5"}],
                 "unit": "USD" if any(w in first[0] for w in ("usd", "savings", "cost")) else "", "source": f"{dataset} from the planner tables"}
        actions.append({"type": "chart", "chart": chart})
    return {**table, "rows": rows[:15], "note": f"{len(rows)} rows in the table card" if len(rows) > 15 else None}, actions


# ---------- registry ----------

def export_tools(ctx):
    filt = {"type": "object", "properties": {"years": {"type": "array", "items": {"type": "integer"}}, "partner": {"type": "string"}, "tier": {"type": "string"},
                                             "verdict": {"type": "string"}, "state": {"type": "string"}, "status": {"type": "string"}, "id": {"type": "string"},
                                             "period": {"type": "string"}, "month": {"type": "integer"}, "days": {"type": "integer"}, "impact": {"type": "string"},
                                             "direction": {"type": "string"}, "top": {"type": "integer"}, "horizon": {"type": "string"}, "starred": {"type": "boolean"}}}
    return {
        "export": (_bind(ctx, export), "Build a file to download: overlaps, projects, plan, findings or hazard_exposure as geojson or kml (for ArcGIS and Google "
                   "Earth), csv, xlsx, or ics (calendar entries for plan target windows or project build windows). Only what we can see.", {
            "kind": {"type": "string", "enum": list(EXPORT_KINDS)}, "format": {"type": "string", "enum": list(FORMATS)}, "filters": filt}, ["kind", "format"]),
        "share_link": (_bind(ctx, share_link), "An expiring link (default 7 days) that shows a report, finding or the latest plan to someone without a login; "
                       "revocable from the card.", {"kind": {"type": "string", "enum": list(share.KINDS)}, "id": {"type": "string", "description": "report or finding id; for plan the horizon"},
                                                    "expires_days": {"type": "integer"}}, ["kind"]),
        "query_data": (_bind(ctx, query_data), "Group, aggregate and sort any dataset (" + ", ".join([*charts.TABLES, "requests", *charts.DATASETS]) + "): e.g. average "
                       "drive_min of overlaps by partner, or sum of savings_high by tier. Returns a table card, plus a chart when grouped.", {
            "dataset": {"type": "string"}, "filters": filt, "group_by": {"type": "array", "items": {"type": "string"}},
            "aggregate": {"type": "object", "description": "field -> sum|avg|min|max|count", "additionalProperties": {"type": "string", "enum": list(AGGS)}},
            "sort": {"type": "object", "properties": {"by": {"type": "string"}, "dir": {"type": "string", "enum": ["asc", "desc"]}}}, "limit": {"type": "integer"}}, ["dataset"]),
    }


PROMPT = """
- "Export as / save as / give me a file / for ArcGIS / for Excel / add to my calendar" means export with the matching format (geojson or kml for GIS,
  xlsx or csv for spreadsheets, ics for calendars). "Share this / send a link / let X see it" means share_link. Questions that group or total data
  ("average drive time by partner", "total savings by tier", "how many projects per state") mean query_data with group_by and aggregate; quote the
  tool's totals only."""
